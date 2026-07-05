import argparse
import json
import subprocess
import sys
from pathlib import Path

REQUIRED_PACKAGES = {
    "fitz": "PyMuPDF",
    "pdfplumber": "pdfplumber",
    "openpyxl": "openpyxl",
    "pandas": "pandas",
}


def ensure_dependencies():
    """Installe automatiquement les packages pip manquants."""
    missing = []
    for import_name, pip_name in REQUIRED_PACKAGES.items():
        try:
            __import__(import_name)
        except ImportError:
            missing.append(pip_name)
    if missing:
        print(f"PROGRESS:install:Installation automatique de {', '.join(missing)}...", flush=True)
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "--quiet"] + missing,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.PIPE,
        )
        print("PROGRESS:install:Dependances document_extract installees.", flush=True)


ensure_dependencies()


def clip(text: str, limit: int = 24000) -> str:
    text = text.replace("\r", "").strip()
    if len(text) <= limit:
        return text
    return text[:limit] + "\n...[contenu tronque pour rester exploitable]"


def render_pdf_previews(document, path: Path, preview_dir: Path | None) -> list[str]:
    if preview_dir is None:
        return []

    import fitz

    preview_dir.mkdir(parents=True, exist_ok=True)
    preview_paths: list[str] = []

    for page_index in range(min(2, len(document))):
        page = document.load_page(page_index)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False)
        preview_path = preview_dir / f"{path.stem}_page_{page_index + 1}.png"
        pixmap.save(preview_path)
        preview_paths.append(str(preview_path))

    return preview_paths


def extract_pdf(path: Path, preview_dir: Path | None = None) -> tuple[str, list[str]]:
    chunks: list[str] = []

    import fitz  # pymupdf

    document = fitz.open(path)
    preview_paths = render_pdf_previews(document, path, preview_dir)
    for page_index, page in enumerate(document, start=1):
        text = page.get_text("text").strip()
        if text:
            chunks.append(f"[Page {page_index}]\n{text}")

    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables() or []
                for table_index, table in enumerate(tables, start=1):
                    rows = []
                    for row in table:
                        safe_row = [cell.strip() if isinstance(cell, str) else "" for cell in row]
                        rows.append(" | ".join(safe_row))
                    if rows:
                        chunks.append(f"[Table page {page_index}.{table_index}]\n" + "\n".join(rows))
    except Exception:
        pass

    return clip("\n\n".join(chunks)), preview_paths


def extract_spreadsheet(path: Path) -> str:
    chunks: list[str] = []

    import pandas as pd
    from openpyxl import load_workbook

    suffix = path.suffix.lower()

    if suffix in {".csv", ".tsv"}:
        separator = "\t" if suffix == ".tsv" else ","
        frame = pd.read_csv(path, sep=separator)
        chunks.append(f"[Sheet main]\n{frame.head(40).to_string(index=False)}")
        return clip("\n\n".join(chunks))

    workbook = load_workbook(path, read_only=True, data_only=True)
    frames = pd.read_excel(path, sheet_name=None)

    for sheet_name in workbook.sheetnames:
        frame = frames.get(sheet_name)
        if frame is not None:
            chunks.append(f"[Sheet {sheet_name}]\n{frame.head(40).fillna('').to_string(index=False)}")
        else:
            worksheet = workbook[sheet_name]
            rows = []
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                safe_row = ["" if cell is None else str(cell) for cell in row]
                rows.append(" | ".join(safe_row))
                if row_index >= 40:
                    break
            if rows:
                chunks.append(f"[Sheet {sheet_name}]\n" + "\n".join(rows))

    return clip("\n\n".join(chunks))


def extract_text(path: Path) -> str:
    return clip(path.read_text(encoding="utf-8", errors="ignore"))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--preview-dir")
    args = parser.parse_args()

    path = Path(args.input)
    suffix = path.suffix.lower()
    preview_dir = Path(args.preview_dir) if args.preview_dir else None

    try:
        if suffix == ".pdf":
            text, preview_paths = extract_pdf(path, preview_dir)
        elif suffix in {".xlsx", ".xls", ".xlsm", ".csv", ".tsv"}:
            text = extract_spreadsheet(path)
            preview_paths = []
        else:
            text = extract_text(path)
            preview_paths = []

        print(json.dumps({"ok": True, "text": text, "preview_paths": preview_paths}, ensure_ascii=False))
        return 0
    except Exception as error:
        print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
